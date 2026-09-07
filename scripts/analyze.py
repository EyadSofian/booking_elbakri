import openpyxl, json, datetime, re, collections, os
from openpyxl.utils import get_column_letter

SHEETS = {
 "BOOKING/Sheet1": ("data/legacy/ELBAKRI OVER SEAS BOOKING .xlsx","Sheet1",8,
   {"A":"name","D":"nationality","F":"phone","H":"checkIn","J":"checkOut","L":"hotel","O":"roomType","Q":"mealPlan","S":"bookingDate","U":"agency","W":"notes","Z":"legacyZ"}),
 "EX/Sheet1": ("data/legacy/ُELBAKRI OVER SEAS EX.xlsx","Sheet1",6,
   {"A":"name","D":"pax","E":"child","G":"phone","I":"nationality","K":"hotel","M":"ex","O":"date","Q":"rest","S":"agency","U":"notes"}),
 "TRANSFER": ("data/legacy/ELBAKRI OVER SEAS FOR TRANSFER .xlsx","TRANSFER",6,
   {"A":"name","D":"phone","F":"from","H":"to","J":"pax","L":"nationality","N":"date","P":"flight","R":"pickup","T":"agency","V":"notes"}),
 "VISA": ("data/legacy/ELBAKRI OVER SEAS FOR TRANSFER .xlsx","VISA",6,
   {"A":"name","D":"phone","F":"from","H":"to","J":"pax","L":"nationality","N":"date","P":"agency","R":"net","T":"sell"}),
 "PAYMENT": ("data/legacy/PYAMNT.xlsx","payment ",3,
   {"A":"hotelName","D":"total","F":"paid","H":"rest","J":"paymentDate","L":"checkIn","N":"status"}),
 "SAMA": ("data/legacy/PYAMNT.xlsx","SAMA",None,
   {"A":"note","D":"amount","H":"c3","J":"c4","L":"c5"}),
}

def norm(v):
    if v is None: return None
    if isinstance(v,str):
        s=v.strip()
        return s if s else None
    return v

def jrepr(v):
    if isinstance(v,(datetime.datetime,datetime.date,datetime.time)): return {"__t":type(v).__name__,"v":v.isoformat()}
    if isinstance(v,datetime.timedelta): return {"__t":"timedelta","v":v.total_seconds()}
    return v

out={}
for key,(path,sheet,header,cols) in SHEETS.items():
    wb=openpyxl.load_workbook(path,data_only=True)
    ws=wb[sheet]
    anchors=set()
    inside=set()
    for rng in ws.merged_cells.ranges:
        anchors.add((rng.min_row,rng.min_col))
        for r in range(rng.min_row,rng.max_row+1):
            for c in range(rng.min_col,rng.max_col+1):
                if (r,c)!=(rng.min_row,rng.min_col): inside.add((r,c))
    colidx={l:openpyxl.utils.column_index_from_string(l) for l in cols}
    recs=[]
    allcols=set()
    for r in range(1,ws.max_row+1):
        row={}
        extra={}
        for c in range(1,ws.max_column+1):
            if (r,c) in inside: continue
            v=norm(ws.cell(row=r,column=c).value)
            if v is None: continue
            L=get_column_letter(c)
            if L in cols: row[cols[L]]=v
            else: extra[L]=v; allcols.add(L)
        if row or extra:
            recs.append({"row":r,"data":{k:jrepr(v) for k,v in row.items()},"extra":{k:jrepr(v) for k,v in extra.items()}})
    out[key]={"path":path,"sheet":sheet,"header":header,"maxRow":ws.max_row,"records":recs,"unmappedCols":sorted(allcols)}
    print(f"{key}: {len(recs)} non-empty logical rows, unmapped cols {sorted(allcols)}")

os.makedirs("docs/_analysis",exist_ok=True)
json.dump(out,open("docs/_analysis/extract.json","w"),ensure_ascii=False,indent=1)
