import openpyxl, sys, json, datetime
from openpyxl.utils import get_column_letter

FILES = [
 ("data/legacy/PYAMNT.xlsx", None),
 ("data/legacy/ELBAKRI OVER SEAS BOOKING .xlsx", None),
 ("data/legacy/ُELBAKRI OVER SEAS EX.xlsx", None),
 ("data/legacy/ELBAKRI OVER SEAS FOR TRANSFER .xlsx", None),
]

def cellrepr(v):
    if v is None: return None
    if isinstance(v, (datetime.datetime, datetime.date, datetime.time)): return f"<{type(v).__name__}:{v.isoformat()}>"
    if isinstance(v, float): return f"<f:{v!r}>"
    if isinstance(v, int): return f"<i:{v}>"
    s = str(v)
    return s

for path, _ in FILES:
    wb = openpyxl.load_workbook(path, data_only=True)
    print("="*100)
    print("FILE:", path)
    print("SHEETS:", wb.sheetnames)
    for ws in wb.worksheets:
        print("-"*100)
        print(f"SHEET: {ws.title!r}  dims={ws.dimensions} max_row={ws.max_row} max_col={ws.max_column}")
        print(f"  merged: {len(ws.merged_cells.ranges)} ranges -> {[str(r) for r in list(ws.merged_cells.ranges)[:20]]}")
        # find non-empty rows
        nonempty = 0
        for r in range(1, min(ws.max_row, 4000)+1):
            vals = [ws.cell(row=r, column=c).value for c in range(1, min(ws.max_column,40)+1)]
            if any(v is not None and str(v).strip()!="" for v in vals):
                nonempty += 1
        print(f"  non-empty rows: {nonempty}")
        # print first 40 rows raw
        for r in range(1, min(ws.max_row, 40)+1):
            vals = []
            for c in range(1, min(ws.max_column,30)+1):
                v = ws.cell(row=r, column=c).value
                if v is not None and str(v).strip()!="":
                    vals.append(f"{get_column_letter(c)}={cellrepr(v)}")
            if vals:
                print(f"   R{r}: " + " | ".join(vals))
